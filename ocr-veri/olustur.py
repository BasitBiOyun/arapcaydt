import os, re, glob, json, zipfile, hashlib, io, shutil, subprocess, sys
from PIL import Image
import docx
sys.stdout.reconfigure(encoding='utf-8')
SRC=os.getcwd(); OUT=os.path.join(os.path.dirname(SRC),'ArapcaYDT_OCR_veri')
os.makedirs(OUT, exist_ok=True)
RLM=re.compile('[‎‏‪-‮]')
def clean(s): return re.sub(r'\s+',' ',RLM.sub('',s)).strip()
# 1) konu testleri: clean jpg + docx text
kt=os.path.join(OUT,'konu_testleri'); os.makedirs(kt+'/images',exist_ok=True); os.makedirs(kt+'/images_cevapli',exist_ok=True)
rows=[]; warn=[]
for d in sorted(glob.glob('ARAPCA_*/*/*')):
    name=os.path.basename(d)
    doc=docx.Document(os.path.join(d,name+'.docx'))
    paras=[p.text for p in doc.paragraphs if p.text.strip()]
    qs=[]; cur=None
    for p in paras:
        m=re.match(r'^\s*(\d+)\s*[-.]\s*(.*)$',p)
        if m and re.match(r'[A-Za-zÇĞİÖŞÜçğıöşüÂâ]',m.group(2)):
            cur={'qnum':int(m.group(1)),'instruction':clean(m.group(2)),'body':[]}; qs.append(cur)
        elif cur: cur['body'].append(p)
    jdir=os.path.join(d,'Soru_ve_Cevaplar_JPG')
    for q in qs:
        body='\n'.join(q['body'])
        parts=re.split(r'(?:^|(?<=\s)|(?<=[؀-ۿ.]))([A-E])\)\s*',body)
        stem=clean(parts[0]); opts={}
        for i in range(1,len(parts)-1,2): opts[parts[i]]=clean(parts[i+1])
        if list(opts)!=list('ABCDE'): warn.append(f'{name} q{q["qnum"]} opts={list(opts)}')
        n=f'{q["qnum"]:02d}'
        img=os.path.join(jdir,f'{name}_{n}.jpg'); ca=os.path.join(jdir,f'{name}_CA_{n}.jpg')
        if not os.path.exists(img): warn.append(f'missing {img}'); continue
        safe=name.replace('Â','A')
        shutil.copy2(img,f'{kt}/images/{safe}_{n}.jpg')
        if os.path.exists(ca): shutil.copy2(ca,f'{kt}/images_cevapli/{safe}_CA_{n}.jpg')
        arabic_text='\n'.join(x for x in [stem]+[opts[k] for k in opts] if x)
        rows.append(dict(image=f'images/{safe}_{n}.jpg',image_cevapli=f'images_cevapli/{safe}_CA_{n}.jpg' if os.path.exists(ca) else None,
            set=safe,qnum=q['qnum'],instruction_tr=q['instruction'],stem_ar=stem,options=opts,arabic_text=arabic_text))
    njpg=len([f for f in os.listdir(jdir) if '_CA_' not in f])
    if njpg!=len(qs): warn.append(f'{name}: docx {len(qs)} soru, jpg {njpg}')
with open(kt+'/labels.jsonl','w',encoding='utf-8') as f:
    for r in rows: f.write(json.dumps(r,ensure_ascii=False)+'\n')
print('konu testleri', len(rows)); print('\n'.join(warn))
# 2) cikmis sorular pptx: largest per-slide image, dedup by hash
cs=os.path.join(OUT,'cikmis_sorular'); seen={}; ccount=0; dup=0
decks=sorted(glob.glob('Çıkmış Sorular Sunu Halleri*/*/*.pptx'))+sorted(glob.glob('çıkmış sorular20*.pptx'))
meta=[]
for p in decks:
    year=re.search(r'(20\d\d)',os.path.basename(p)).group(1)
    z=zipfile.ZipFile(p); names=set(z.namelist())
    slides=sorted([n for n in names if re.match(r'ppt/slides/slide\d+\.xml$',n)],key=lambda n:int(re.findall(r'\d+',n)[-1]))
    for s in slides:
        sn=int(re.findall(r'\d+',s)[-1]); rel=f'ppt/slides/_rels/slide{sn}.xml.rels'
        if rel not in names: continue
        media=['ppt/'+m for m in re.findall(r'\.\./(media/[^"]+\.(?:png|jpe?g))',z.read(rel).decode(),re.I)]
        best=None
        for m in media:
            b=z.read(m); im=Image.open(io.BytesIO(b))
            if im.size[1]<=130: continue  # banners/headers
            if best is None or im.size[0]*im.size[1]>best[1]: best=(b,im.size[0]*im.size[1],m,im.size)
        if not best: continue
        h=hashlib.sha1(best[0]).hexdigest()
        if h in seen: dup+=1; continue
        ext=os.path.splitext(best[2])[1].lower()
        rp=f'{year}/slide_{sn:02d}{ext}'; os.makedirs(os.path.join(cs,year),exist_ok=True)
        open(os.path.join(cs,rp),'wb').write(best[0]); seen[h]=rp; ccount+=1
        meta.append(dict(image=rp,year=int(year),slide=sn,size=best[3],source=os.path.relpath(p,SRC),arabic_text=None))
with open(cs+'/images.jsonl','w',encoding='utf-8') as f:
    for r in meta: f.write(json.dumps(r,ensure_ascii=False)+'\n')
print('cikmis', ccount,'dup skipped',dup)
# 3) 2026 slides png
s26=os.path.join(OUT,'slaytlar_2026_1920x1080'); os.makedirs(s26,exist_ok=True)
for p in glob.glob('YDT_Arabca_Sorular_1920x1080/*/*.png'): shutil.copy2(p,s26)
# 4) deneme1: pdf + per-page text + kaynak
dn=os.path.join(OUT,'deneme1'); os.makedirs(dn+'/sayfa_metin',exist_ok=True)
pdf='YDT_Arapca_Deneme_1_v2.pdf'; shutil.copy2(pdf,dn)
t=subprocess.run(['pdftotext','-enc','UTF-8',pdf,'-'],capture_output=True).stdout.decode('utf8')
pages=t.split('\f')
for i,pg in enumerate(pages,1):
    if pg.strip(): open(f'{dn}/sayfa_metin/sayfa_{i:02d}.txt','w',encoding='utf-8').write(RLM.sub('',pg))
for p in glob.glob('deneme1_kaynak_v2/data*.py'): shutil.copy2(p,dn)
# 5) analiz docx -> txt corpus
an=os.path.join(OUT,'metin_korpusu'); os.makedirs(an,exist_ok=True); na=0
for p in glob.glob('**/*.docx',recursive=True):
    b=os.path.basename(p)
    if b.startswith('~$') or b.startswith('ARAPCA_') or 'Deneme' in b: continue
    try: txt='\n'.join(x.text for x in docx.Document(p).paragraphs)
    except Exception as e: print('skip',p,e); continue
    open(os.path.join(an,os.path.splitext(b)[0]+'.txt'),'w',encoding='utf-8').write(txt); na+=1
print('korpus', na, 'deneme pages', len([x for x in pages if x.strip()]))
